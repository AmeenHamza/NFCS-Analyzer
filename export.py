"""Export the proposal to a Word document."""

import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"


def _set_font(style_or_run, east_asian: str, size: float = None, bold: bool = None) -> None:
    font = style_or_run.font
    font.name = LATIN_FONT
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    element = style_or_run.element
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asian)


def _add_inline(paragraph, text: str) -> None:
    """Add text, turning **bold** spans into bold runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            run = paragraph.add_run(part)
            if i % 2 == 1:
                run.bold = True


def _add_table(document, rows: list[str]) -> None:
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    table = document.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(cells):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = ""
            _add_inline(cell.paragraphs[0], row[c] if c < len(row) else "")
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph()


def _add_body(document, text: str) -> None:
    """Render the light Markdown the prompts ask for: ###, -, 1., tables, **bold**."""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            _add_table(document, block)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            document.add_heading(heading.group(2).strip("* "), level=max(level, 2))
        elif re.match(r"^[-*•]\s+", stripped):
            _add_inline(document.add_paragraph(style="List Bullet"),
                        re.sub(r"^[-*•]\s+", "", stripped))
        elif re.match(r"^\d+(?:[.)]\s+|、)", stripped):
            _add_inline(document.add_paragraph(style="List Number"),
                        re.sub(r"^\d+(?:[.)]\s+|、)", "", stripped))
        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Pt(24)
            _add_inline(paragraph, stripped)
        i += 1


def to_docx(title: str, sections, subtitle: str = "") -> bytes:
    """Build the .docx: title page, then a heading and body for each section.

    `sections` is a list of (heading, text) pairs or a {heading: text} dict.
    """
    if isinstance(sections, dict):
        sections = list(sections.items())

    document = Document()
    normal = document.styles["Normal"]
    _set_font(normal, BODY_FONT, size=12)
    normal.paragraph_format.line_spacing = 1.5
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        _set_font(document.styles[name], HEADING_FONT)

    # Title page
    for _ in range(6):
        document.add_paragraph()
    if subtitle:
        p = document.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.runs[0], HEADING_FONT, size=16)
    p = document.add_paragraph("项 目 申 报 书")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.runs[0], HEADING_FONT, size=26, bold=True)
    document.add_paragraph()
    p = document.add_paragraph(title)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.runs[0], HEADING_FONT, size=18, bold=True)
    for _ in range(8):
        document.add_paragraph()
    p = document.add_paragraph(f"申请日期：{date.today():%Y年%m月%d日}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].add_break(WD_BREAK.PAGE)

    for number, (heading, body) in enumerate(sections, start=1):
        document.add_heading(f"{_cn_number(number)}、{heading}", level=1)
        _add_body(document, body or "")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _cn_number(n: int) -> str:
    digits = "零一二三四五六七八九十"
    if n <= 10:
        return digits[n]
    if n < 20:
        return "十" + digits[n - 10]
    return digits[n // 10] + "十" + (digits[n % 10] if n % 10 else "")
